"""The admin punch log: list, summary, detail, photo view, photo comparison and Excel export,
with every role (spec M2 §4.10, §5.3, §7 rules 1, 17, 21; AC-23, AC-24, AC-33, AC-34, AC-37).
Real PostgreSQL and a real local photo folder; fake data only."""

import io
import json
import logging
import uuid
from datetime import UTC, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.crypto import decode_key, encrypt_snapshot
from app.core.storage import StorageError, get_storage
from app.core.time import IST, today_ist, utc_now
from app.modules.audit.models import AuditLog
from app.modules.employees.models import Employee
from app.modules.kiosks.models import KioskDevice
from app.modules.punches import admin_service
from app.modules.punches.models import Punch, PunchSnapshot
from app.modules.users.models import Role, User
from app.modules.users.plant_access import set_user_plant
from tests.conftest import clear_cached_settings
from tests.image_helpers import make_jpeg
from tests.people_helpers import PlantIds, client_for, create_plant, temporary_body

PUNCHES = "/api/v1/punches"


@pytest.fixture(autouse=True)
def photo_folder(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("STORAGE_LOCAL_DIR", str(tmp_path / "photos"))
    clear_cached_settings()


@pytest.fixture
def boss(db_app: FastAPI, db_session: Session) -> TestClient:
    return client_for(db_app, db_session, Role.SUPER_ADMIN, "boss")


@pytest.fixture
def k134(boss: TestClient, db_session: Session) -> PlantIds:
    return create_plant(boss, db_session, "K134")


@pytest.fixture
def tp1(boss: TestClient, db_session: Session) -> PlantIds:
    return create_plant(boss, db_session, "TP1")


@pytest.fixture
def corporate(db_app: FastAPI, db_session: Session) -> TestClient:
    return client_for(db_app, db_session, Role.CORPORATE_HEAD, "corp")


def head_of(db_app: FastAPI, db_session: Session, plant: PlantIds, username: str) -> TestClient:
    head = client_for(db_app, db_session, Role.PLANT_HEAD, username)
    user = db_session.scalars(select(User).where(User.username == username)).one()
    set_user_plant(db_session, user, uuid.UUID(plant.id))
    db_session.commit()
    return head


def make_kiosk(session: Session, plant: PlantIds, name: str) -> KioskDevice:
    kiosk = KioskDevice(plant_id=uuid.UUID(plant.id), name=name, status="active")
    session.add(kiosk)
    session.commit()
    return kiosk


class Card:
    def __init__(self, boss: TestClient, plant: PlantIds, code: str, name: str) -> None:
        body = temporary_body(plant) | {"employee_code": code, "full_name": name}
        employee = boss.post("/api/v1/employees", json=body).json()
        card = boss.post(f"/api/v1/employees/{employee['id']}/qr-cards", json={}).json()["card"]
        self.employee_id = uuid.UUID(employee["id"])
        self.card_id = uuid.UUID(card["id"])


@pytest.fixture
def asha(boss: TestClient, k134: PlantIds) -> Card:
    return Card(boss, k134, "T-2041", "Asha Kale")


@pytest.fixture
def ravi(boss: TestClient, tp1: PlantIds) -> Card:
    return Card(boss, tp1, "T-3050", "Ravi Shinde")


def at_ist(hour: int, minute: int = 0, days_ago: int = 0) -> datetime:
    day = today_ist() - timedelta(days=days_ago)
    return datetime(day.year, day.month, day.day, hour, minute, tzinfo=IST)


def add_punch(
    session: Session,
    kiosk: KioskDevice,
    card: Card | None,
    when: datetime,
    *,
    direction: str | None = "in",
    reason: str | None = None,
) -> Punch:
    punch = Punch(
        id=uuid.uuid4(),
        client_punch_id=uuid.uuid4(),
        kiosk_id=kiosk.id,
        plant_id=kiosk.plant_id,
        employee_id=card.employee_id if card else None,
        qr_credential_id=card.card_id if card else None,
        method="qr",
        direction=None if reason else direction,
        decision="rejected" if reason else "accepted",
        reason_codes=[reason] if reason else [],
        server_received_at=when,
        device_claimed_at=when - timedelta(seconds=3),
        lat=19.85,
        lon=75.2,
        accuracy_m=12,
        ip="203.0.113.7",
        app_version="abc123",
    )
    session.add(punch)
    session.commit()
    return punch


def add_photo(
    session: Session,
    punch: Punch,
    *,
    state: str = "stored",
    is_dark: bool = False,
    write_file: bool = True,
) -> PunchSnapshot:
    key = f"snapshots/2026/10/20/{punch.id}.bin"
    if state == "stored" and write_file:
        secret = get_settings().snapshot_encryption_key
        assert secret is not None
        jpeg = make_jpeg(colour=(10, 10, 10) if is_dark else (150, 150, 150))
        get_storage().save(
            key, encrypt_snapshot(jpeg, punch.id, decode_key(secret.get_secret_value()))
        )
    now = utc_now()
    snapshot = PunchSnapshot(
        punch_id=punch.id,
        storage_key=key,
        size_bytes=100,
        width=320,
        height=240,
        brightness=10 if is_dark else 150,
        is_dark=is_dark,
        taken_at=punch.server_received_at,
        expires_at=now + timedelta(days=60),
        state=state,
        purged_at=now if state == "purged" else None,
    )
    session.add(snapshot)
    session.commit()
    return snapshot


@pytest.fixture
def day(db_session: Session, k134: PlantIds, tp1: PlantIds, asha: Card, ravi: Card) -> dict:
    """Today at two plants: K134 has an IN, an OUT, a blocked card, a visiting card (WRONG_PLANT)
    and an unknown card; TP1 has one IN. Yesterday late evening K134 had one IN."""
    main_gate = make_kiosk(db_session, k134, "K134 main gate")
    tp1_gate = make_kiosk(db_session, tp1, "TP1 gate")
    punches = {
        "in": add_punch(db_session, main_gate, asha, at_ist(7, 2)),
        "out": add_punch(db_session, main_gate, asha, at_ist(15, 30), direction="out"),
        "revoked": add_punch(db_session, main_gate, asha, at_ist(9, 0), reason="QR_REVOKED"),
        "visitor": add_punch(db_session, main_gate, ravi, at_ist(10, 0), reason="WRONG_PLANT"),
        "unknown": add_punch(
            db_session, main_gate, None, at_ist(11, 0), reason="QR_SIGNATURE_INVALID"
        ),
        "tp1": add_punch(db_session, tp1_gate, ravi, at_ist(8, 0)),
        "yesterday": add_punch(db_session, main_gate, asha, at_ist(23, 30, days_ago=1)),
    }
    add_photo(db_session, punches["in"])
    add_photo(db_session, punches["revoked"], is_dark=True)
    add_photo(db_session, punches["visitor"], state="purged")
    add_photo(db_session, punches["unknown"], state="upload_failed")
    return {"punches": punches, "main_gate": main_gate, "tp1_gate": tp1_gate}


def ids(response) -> set[str]:
    assert response.status_code == 200, response.text
    return {item["id"] for item in response.json()["items"]}


def pid(day: dict, name: str) -> str:
    return str(day["punches"][name].id)


# --- The list -----------------------------------------------------------------------------------


def test_the_super_admin_sees_todays_punches_of_every_plant_newest_first(
    boss: TestClient, day: dict
) -> None:
    page = boss.get(PUNCHES).json()

    names = ["out", "unknown", "visitor", "revoked", "tp1", "in"]
    assert [item["id"] for item in page["items"]] == [pid(day, name) for name in names]
    assert page["total"] == 6  # yesterday's 23:30 IST punch is another day


def test_a_row_names_the_person_kiosk_plant_result_and_photo_state(
    boss: TestClient, day: dict
) -> None:
    items = {item["id"]: item for item in boss.get(PUNCHES).json()["items"]}

    first = items[pid(day, "in")]
    assert first["employee"]["code"] == "T-2041" and first["employee"]["name"] == "Asha Kale"
    assert first["kiosk"]["name"] == "K134 main gate" and first["plant"]["code"] == "K134"
    assert (first["direction"], first["decision"], first["reason_codes"]) == ("in", "accepted", [])
    assert first["snapshot"]["state"] == "stored" and first["snapshot"]["is_dark"] is False
    assert items[pid(day, "revoked")]["snapshot"]["is_dark"] is True
    assert items[pid(day, "visitor")]["snapshot"]["state"] == "purged"
    assert items[pid(day, "unknown")]["snapshot"]["state"] == "upload_failed"
    assert items[pid(day, "unknown")]["employee"] is None
    assert items[pid(day, "out")]["snapshot"] == {
        "state": "none",
        "is_dark": None,
        "expires_at": None,
        "purged_at": None,
        "face_found": None,  # M3 F-30: no photo, so nothing to check
    }


def test_no_answer_contains_the_ip_location_or_storage_key(boss: TestClient, day: dict) -> None:
    text = json.dumps(boss.get(PUNCHES).json())

    assert "203.0.113.7" not in text
    assert "storage_key" not in text and "snapshots/" not in text
    assert "19.85" not in text


def test_filters_by_day_plant_result_kiosk_reason_and_search(
    boss: TestClient, day: dict, k134: PlantIds, asha: Card
) -> None:
    yesterday = (today_ist() - timedelta(days=1)).isoformat()
    assert ids(boss.get(PUNCHES, params={"date": yesterday})) == {pid(day, "yesterday")}
    assert pid(day, "tp1") not in ids(boss.get(PUNCHES, params={"plant_id": k134.id}))
    assert ids(boss.get(PUNCHES, params={"decision": "accepted"})) == {
        pid(day, "in"),
        pid(day, "out"),
        pid(day, "tp1"),
    }
    assert ids(boss.get(PUNCHES, params={"kiosk_id": str(day["tp1_gate"].id)})) == {pid(day, "tp1")}
    assert ids(boss.get(PUNCHES, params={"reason_code": "WRONG_PLANT"})) == {pid(day, "visitor")}
    assert ids(boss.get(PUNCHES, params={"employee_id": str(asha.employee_id)})) == {
        pid(day, "in"),
        pid(day, "out"),
        pid(day, "revoked"),
    }
    assert ids(boss.get(PUNCHES, params={"q": "t-30"})) == {pid(day, "visitor"), pid(day, "tp1")}
    assert ids(boss.get(PUNCHES, params={"q": "kale"})) == {
        pid(day, "in"),
        pid(day, "out"),
        pid(day, "revoked"),
    }
    assert ids(boss.get(PUNCHES, params={"dark_only": "true"})) == {pid(day, "revoked")}


def test_the_corporate_head_reads_every_plant_but_never_learns_about_dark_photos(
    corporate: TestClient, day: dict
) -> None:
    items = corporate.get(PUNCHES).json()["items"]

    assert len(items) == 6
    assert all(item["snapshot"]["is_dark"] is None for item in items)
    # "Dark photos only" is a Super Admin filter: ignored for anyone else.
    assert len(corporate.get(PUNCHES, params={"dark_only": "true"}).json()["items"]) == 6


def test_a_plant_head_sees_only_punches_at_their_plants_kiosks(
    db_app: FastAPI, db_session: Session, day: dict, k134: PlantIds, tp1: PlantIds
) -> None:
    head = head_of(db_app, db_session, k134, "head.k134")

    assert pid(day, "tp1") not in ids(head.get(PUNCHES))
    assert len(ids(head.get(PUNCHES))) == 5
    assert head.get(PUNCHES, params={"plant_id": tp1.id}).json()["error"]["code"] == "NOT_FOUND"
    assert head.get(f"{PUNCHES}/{pid(day, 'tp1')}").json()["error"]["code"] == "NOT_FOUND"


def test_a_plant_head_never_sees_who_a_card_of_another_plant_belongs_to(
    db_app: FastAPI, db_session: Session, boss: TestClient, day: dict, k134: PlantIds
) -> None:
    head = head_of(db_app, db_session, k134, "head.k134")

    visitor = next(i for i in head.get(PUNCHES).json()["items"] if i["id"] == pid(day, "visitor"))
    seen_by_boss = next(
        i for i in boss.get(PUNCHES).json()["items"] if i["id"] == pid(day, "visitor")
    )

    assert visitor["employee"] is None and visitor["employee_hidden"] is True
    assert "Ravi" not in json.dumps(head.get(PUNCHES).json())
    assert "T-3050" not in json.dumps(head.get(PUNCHES).json())
    assert seen_by_boss["employee"]["name"] == "Ravi Shinde"
    assert seen_by_boss["employee_hidden"] is False


def test_the_punch_log_needs_a_login(db_client: TestClient) -> None:
    assert db_client.get(PUNCHES).status_code == 401


# --- Summary and detail -------------------------------------------------------------------------


def test_the_day_summary_counts_results_and_reasons(
    boss: TestClient, day: dict, k134: PlantIds
) -> None:
    summary = boss.get(f"{PUNCHES}/summary", params={"plant_id": k134.id}).json()

    assert summary == {
        "accepted_in": 1,
        "accepted_out": 1,
        "rejected": 3,
        "dark_photos": 1,
        "rejected_by_reason": {"QR_REVOKED": 1, "WRONG_PLANT": 1, "QR_SIGNATURE_INVALID": 1},
        # M3: face punches and the face-in-photo marks of card photos.
        "face_accepted": 0,
        "code_accepted": 0,
        "asked_for_code": 0,
        "no_face_photos": 0,
    }
    everywhere = boss.get(f"{PUNCHES}/summary").json()
    assert everywhere["accepted_in"] == 2


def test_only_the_super_admin_gets_the_dark_photo_count(
    db_app: FastAPI, db_session: Session, corporate: TestClient, day: dict, k134: PlantIds
) -> None:
    assert corporate.get(f"{PUNCHES}/summary").json()["dark_photos"] is None
    head = head_of(db_app, db_session, k134, "head.k134")
    summary = head.get(f"{PUNCHES}/summary").json()
    assert summary["dark_photos"] is None and summary["accepted_in"] == 1


def test_the_detail_adds_version_device_time_and_location(boss: TestClient, day: dict) -> None:
    detail = boss.get(f"{PUNCHES}/{pid(day, 'in')}").json()

    assert detail["app_version"] == "abc123"
    assert detail["device_time"] is not None
    assert detail["location"] == {"lat": 19.85, "lon": 75.2, "accuracy_m": 12.0}
    assert "ip" not in detail
    assert boss.get(f"{PUNCHES}/{uuid.uuid4()}").json()["error"]["code"] == "NOT_FOUND"


# --- Photos -------------------------------------------------------------------------------------


def views(session: Session) -> list[AuditLog]:
    return list(session.scalars(select(AuditLog).where(AuditLog.action == "snapshot.view")))


def test_the_super_admin_opens_a_photo_and_every_view_is_audited(
    boss: TestClient, db_session: Session, day: dict
) -> None:
    first = boss.get(f"{PUNCHES}/{pid(day, 'in')}/snapshot")
    second = boss.get(f"{PUNCHES}/{pid(day, 'in')}/snapshot")

    assert first.status_code == 200 and first.content[:2] == b"\xff\xd8"
    assert first.headers["content-type"] == "image/jpeg"
    assert first.headers["cache-control"] == "no-store"
    assert first.headers["content-disposition"] == "inline"
    assert second.status_code == 200
    entries = views(db_session)
    assert len(entries) == 2
    assert entries[0].entity_id == day["punches"]["in"].id
    assert entries[0].plant_id == day["punches"]["in"].plant_id


def test_other_roles_cannot_open_photos_or_compare(
    db_app: FastAPI, db_session: Session, corporate: TestClient, day: dict, k134: PlantIds
) -> None:
    head = head_of(db_app, db_session, k134, "head.k134")
    for client in (corporate, head):
        for suffix in ("snapshot", "compare"):
            response = client.get(f"{PUNCHES}/{pid(day, 'in')}/{suffix}")
            assert response.json()["error"]["code"] == "FORBIDDEN"
    assert views(db_session) == []


@pytest.mark.parametrize(
    ("name", "status", "code"),
    [
        ("visitor", 410, "SNAPSHOT_PURGED"),
        ("unknown", 404, "SNAPSHOT_NOT_AVAILABLE"),
        ("out", 404, "SNAPSHOT_NOT_AVAILABLE"),
    ],
)
def test_a_photo_that_is_gone_or_was_never_saved_says_so(
    boss: TestClient, db_session: Session, day: dict, name: str, status: int, code: str
) -> None:
    response = boss.get(f"{PUNCHES}/{pid(day, name)}/snapshot")

    assert response.status_code == status
    assert response.json()["error"]["code"] == code
    assert views(db_session) == []


def test_a_missing_file_or_a_file_of_another_punch_is_not_shown(
    boss: TestClient, db_session: Session, day: dict, caplog: pytest.LogCaptureFixture
) -> None:
    main_gate = day["main_gate"]
    no_file = add_punch(db_session, main_gate, None, at_ist(12, 0), reason="QR_FORMAT_INVALID")
    add_photo(db_session, no_file, write_file=False)
    swapped = add_punch(db_session, main_gate, None, at_ist(12, 5), reason="QR_FORMAT_INVALID")
    swapped_photo = add_photo(db_session, swapped, write_file=False)
    # Copy another punch's encrypted file onto this one: the punch id no longer matches.
    in_photo = db_session.scalars(
        select(PunchSnapshot).where(PunchSnapshot.punch_id == day["punches"]["in"].id)
    ).one()
    get_storage().save(swapped_photo.storage_key, get_storage().read(in_photo.storage_key))

    with caplog.at_level(logging.ERROR):
        missing = boss.get(f"{PUNCHES}/{no_file.id}/snapshot")
        moved = boss.get(f"{PUNCHES}/{swapped.id}/snapshot")

    assert missing.json()["error"]["code"] == "SNAPSHOT_NOT_AVAILABLE"
    assert moved.json()["error"]["code"] == "SNAPSHOT_NOT_AVAILABLE"
    assert "could not be decrypted" in caplog.text
    assert views(db_session) == []


def test_a_storage_outage_is_reported_as_unavailable(
    boss: TestClient, day: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    class Down:
        def read(self, key: str) -> bytes:
            raise StorageError("down")

    monkeypatch.setattr(admin_service, "get_storage", Down)

    response = boss.get(f"{PUNCHES}/{pid(day, 'in')}/snapshot")

    assert response.status_code == 503
    assert response.json()["error"]["code"] == "STORAGE_UNAVAILABLE"


def test_compare_shows_the_same_cards_three_newest_earlier_photos(
    boss: TestClient, db_session: Session, day: dict, asha: Card
) -> None:
    gate = day["main_gate"]
    earlier = [
        add_punch(db_session, gate, asha, at_ist(6, minute, days_ago=1)) for minute in (1, 2, 3, 4)
    ]
    for punch in earlier:
        add_photo(db_session, punch, write_file=False)
    gone = add_punch(db_session, gate, asha, at_ist(6, 30, days_ago=1))
    add_photo(db_session, gone, state="purged")
    latest = add_punch(db_session, gate, asha, at_ist(16, 0), direction="out")

    answer = boss.get(f"{PUNCHES}/{latest.id}/compare").json()

    assert answer["current"]["id"] == str(latest.id)
    # Today's 07:02 IN and 09:00 rejection have photos too; then the newest of yesterday.
    expected = [pid(day, "revoked"), pid(day, "in"), str(earlier[3].id)]
    assert [item["id"] for item in answer["earlier"]] == expected
    assert views(db_session) == []  # the list itself shows no photo


def test_compare_for_an_unknown_card_has_no_earlier_photos(boss: TestClient, day: dict) -> None:
    answer = boss.get(f"{PUNCHES}/{pid(day, 'unknown')}/compare").json()

    assert answer["earlier"] == []


# --- Excel export -------------------------------------------------------------------------------


def open_sheet(response) -> list[list[str | None]]:
    assert response.status_code == 200, response.text
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    return [list(row) for row in book.active.iter_rows(values_only=True)]  # type: ignore[union-attr]


def test_the_export_has_the_days_punches_as_text_and_is_audited(
    boss: TestClient, db_session: Session, day: dict, k134: PlantIds
) -> None:
    response = boss.get(f"{PUNCHES}/export.xlsx", params={"plant_id": k134.id})

    rows = open_sheet(response)
    assert rows[0] == [
        "Date",
        "Time (IST)",
        "Plant (kiosk)",
        "Employee code",
        "Name",
        "Kiosk",
        "Direction",
        "Result",
        "Reason code",
        "Reason",
        "Photo",
    ]
    assert len(rows) == 6  # header + 5 punches at K134 today
    first_in = next(row for row in rows if row[1] == "07:02:00")
    assert first_in[3:8] == ["T-2041", "Asha Kale", "K134 main gate", "IN", "Accepted"]
    assert first_in[10] == "Photo stored"
    revoked = next(row for row in rows if row[8] == "QR_REVOKED")
    assert revoked[9] == "Card blocked" and revoked[7] == "Rejected"
    unknown = next(row for row in rows if row[8] == "QR_SIGNATURE_INVALID")
    assert unknown[4] == "Unknown card" and unknown[10] == "Photo not saved"
    assert str(next(row for row in rows if row[8] == "WRONG_PLANT")[10]).startswith(
        "Photo deleted on "
    )
    assert "203.0.113.7" not in str(rows)
    assert response.headers["content-disposition"] == (
        f'attachment; filename="punches-K134-{today_ist():%Y-%m-%d}.xlsx"'
    )
    entry = db_session.scalars(select(AuditLog).where(AuditLog.action == "punches.export")).one()
    assert entry.after is not None and entry.after["rows"] == 5
    assert entry.after["filters"]["plant_id"] == k134.id


def test_a_formula_in_a_name_stays_plain_text(
    boss: TestClient, db_session: Session, day: dict, asha: Card
) -> None:
    employee = db_session.get(Employee, asha.employee_id)
    assert employee is not None
    employee.full_name = '=HYPERLINK("http://example.invalid","click")'
    db_session.commit()

    response = boss.get(f"{PUNCHES}/export.xlsx")

    book = openpyxl.load_workbook(io.BytesIO(response.content))
    names = [cell.value for cell in book.active["E"]]  # type: ignore[index]
    cell = next(c for c in book.active["E"] if c.value and str(c.value).startswith("="))  # type: ignore[index]
    assert '=HYPERLINK("http://example.invalid","click")' in names
    assert cell.data_type == "s"  # a string, not a formula


def test_a_plant_heads_export_hides_people_of_other_plants(
    db_app: FastAPI, db_session: Session, day: dict, k134: PlantIds
) -> None:
    head = head_of(db_app, db_session, k134, "head.k134")

    rows = open_sheet(head.get(f"{PUNCHES}/export.xlsx"))

    assert len(rows) == 6
    assert "Ravi Shinde" not in str(rows) and "T-3050" not in str(rows)
    assert ["", "Card of another plant"] in [row[3:5] for row in rows]


def test_a_day_with_too_many_punches_is_refused(
    boss: TestClient, day: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(admin_service, "MAX_EXPORT_ROWS", 2)

    response = boss.get(f"{PUNCHES}/export.xlsx")

    assert response.json()["error"]["code"] == "EXPORT_TOO_LARGE"


def test_the_export_follows_the_list_filters(boss: TestClient, day: dict) -> None:
    rows = open_sheet(boss.get(f"{PUNCHES}/export.xlsx", params={"decision": "accepted"}))

    assert [row[7] for row in rows[1:]] == ["Accepted"] * 3


def test_times_are_shown_in_ist(boss: TestClient, db_session: Session, day: dict) -> None:
    # 07:02 IST is 01:32 UTC: the export shows the plant's local time.
    assert day["punches"]["in"].server_received_at.astimezone(UTC).hour == 1
    rows = open_sheet(boss.get(f"{PUNCHES}/export.xlsx"))
    assert any(row[1] == "07:02:00" for row in rows)


# --- Security review of M2: fixes ---------------------------------------------------------------


def test_a_plant_head_cannot_find_people_of_other_plants_by_search_or_id(
    db_app: FastAPI, db_session: Session, day: dict, k134: PlantIds, ravi: Card, asha: Card
) -> None:
    """Ravi (plant TP1) scanned at the K134 gate. The row is shown as "card of another plant",
    so a search or an id filter must not confirm who it is (letter by letter)."""
    head = head_of(db_app, db_session, k134, "head.k134")

    by_name = head.get(PUNCHES, params={"q": "Ravi"}).json()["items"]
    by_code_start = head.get(PUNCHES, params={"q": "T-30"}).json()["items"]
    by_id = head.get(PUNCHES, params={"employee_id": str(ravi.employee_id)}).json()["items"]

    assert by_name == [] and by_code_start == [] and by_id == []
    # Their own plant's people are still found.
    assert head.get(PUNCHES, params={"q": "Asha"}).json()["items"] != []
    assert head.get(PUNCHES, params={"employee_id": str(asha.employee_id)}).json()["items"] != []


def test_the_photo_view_entry_says_whose_photo_it_was(
    boss: TestClient, db_session: Session, day: dict
) -> None:
    boss.get(f"{PUNCHES}/{pid(day, 'in')}/snapshot")

    (entry,) = views(db_session)
    assert entry.after is not None
    assert entry.after["employee_id"] == str(day["punches"]["in"].employee_id)

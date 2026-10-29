"""M4 "Needs review": the lists, the review panel, confirm / not accepted / approve, the digest and
the repeat-pattern hints, and the punch log's review marks for every role (spec M4 §4.8, §4.9,
§5.3, §5.4, §7 rules 6, 7, 20; AC-24 to AC-29, AC-42, AC-43). Real PostgreSQL, fake data only."""

import io
import uuid
from datetime import datetime, time, timedelta
from pathlib import Path

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.time import IST, to_ist, today_ist, utc_now
from app.modules.audit.models import AuditLog
from app.modules.kiosks.models import KioskDevice, KioskHeartbeat
from app.modules.org.models import Geofence
from app.modules.punches.models import Punch, PunchReview
from app.modules.users.models import Role, User
from app.modules.users.plant_access import set_user_plant
from tests.conftest import clear_cached_settings
from tests.kiosk_helpers import GATE_LAT, GATE_LON, Device, make_active_kiosk
from tests.people_helpers import PlantIds, client_for, create_plant
from tests.test_punch_api import Worker, make_worker, send

REVIEWS = "/api/v1/reviews"
PUNCHES = "/api/v1/punches"


@pytest.fixture(autouse=True)
def snapshot_folder(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("STORAGE_LOCAL_DIR", str(tmp_path / "snapshots"))
    clear_cached_settings()


@pytest.fixture
def boss(db_app: FastAPI, db_session: Session) -> TestClient:
    return client_for(db_app, db_session, Role.SUPER_ADMIN, "boss")


@pytest.fixture
def k134(boss: TestClient, db_session: Session) -> PlantIds:
    return create_plant(boss, db_session, "K134")


@pytest.fixture
def device() -> Device:
    return Device()


@pytest.fixture
def kiosk(db_session: Session, k134: PlantIds, device: Device) -> KioskDevice:
    return make_active_kiosk(db_session, k134, device)


@pytest.fixture
def worker(boss: TestClient, db_session: Session, k134: PlantIds) -> Worker:
    return make_worker(boss, db_session, k134)


@pytest.fixture
def corporate(db_app: FastAPI, db_session: Session) -> TestClient:
    return client_for(db_app, db_session, Role.CORPORATE_HEAD, "corp")


@pytest.fixture
def plant_head(db_app: FastAPI, db_session: Session, k134: PlantIds) -> TestClient:
    head = client_for(db_app, db_session, Role.PLANT_HEAD, "head")
    user = db_session.scalars(select(User).where(User.username == "head")).one()
    set_user_plant(db_session, user, uuid.UUID(k134.id))
    db_session.commit()
    return head


def add_punch(
    session: Session,
    kiosk: KioskDevice,
    worker: Worker | None,
    *,
    age: timedelta,
    decision: str = "flagged",
    direction: str | None = "in",
    codes: tuple[str, ...] = ("LOCATION_AT_EDGE",),
    method: str = "qr",
    claimed: Worker | None = None,
    offline_delay: timedelta | None = None,
    distance_m: float | None = None,
    lat: float | None = None,
    lon: float | None = None,
) -> Punch:
    """A punch row as the engine would have stored it (faster than going through the gate)."""
    received = utc_now() - age
    is_face = method == "face"
    punch = Punch(
        client_punch_id=uuid.uuid4(),
        kiosk_id=kiosk.id,
        plant_id=kiosk.plant_id,
        employee_id=worker.employee_id if worker else None,
        claimed_employee_id=claimed.employee_id if claimed else None,
        qr_credential_id=worker.card_id if worker and not is_face else None,
        method=method,
        match_mode="code" if is_face and claimed else None,
        direction=direction if decision in ("accepted", "flagged") else None,
        decision=decision,
        reason_codes=list(codes) if decision != "accepted" else [],
        server_received_at=received,
        was_offline=offline_delay is not None,
        offline_punch_time=received - offline_delay if offline_delay else None,
        time_proof="anchored" if offline_delay else None,
        geofence_distance_m=distance_m,
        lat=lat,
        lon=lon,
    )
    session.add(punch)
    session.commit()
    return punch


def review(
    client: TestClient,
    punch: Punch,
    action: str,
    reason: str | None = None,
    direction: str | None = None,
    expected: str | None = None,
):
    body: dict[str, object] = {"action": action, "expected_latest_review_id": expected}
    if reason is not None:
        body["reason"] = reason
    if direction is not None:
        body["direction"] = direction
    return client.post(f"{PUNCHES}/{punch.id}/reviews", json=body)


def ids(response) -> list[str]:
    assert response.status_code == 200, response.text
    return [item["punch"]["id"] for item in response.json()["items"]]


def audit_actions(session: Session) -> list[str]:
    return [
        entry.action
        for entry in session.scalars(
            select(AuditLog).where(AuditLog.action.like("punch_review.%")).order_by(AuditLog.at)
        )
    ]


# --- The lists and the badge (AC-24) -------------------------------------------------------------


def test_the_tabs_list_the_right_punches_and_the_badge_equals_to_check(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    edge = add_punch(db_session, kiosk, worker, age=timedelta(minutes=50))
    travel = add_punch(
        db_session,
        kiosk,
        worker,
        age=timedelta(minutes=40),
        direction="out",
        codes=("IMPOSSIBLE_TRAVEL",),
        offline_delay=timedelta(minutes=52),
    )
    add_punch(db_session, kiosk, worker, age=timedelta(minutes=60), decision="accepted")
    expired = add_punch(
        db_session,
        kiosk,
        worker,
        age=timedelta(minutes=30),
        decision="rejected",
        codes=("WRONG_PLANT",),
    )
    repeat = add_punch(
        db_session,
        kiosk,
        worker,
        age=timedelta(minutes=20),
        decision="rejected",
        codes=("COOLDOWN_ACTIVE",),
    )
    unknown = add_punch(
        db_session,
        kiosk,
        None,
        age=timedelta(minutes=10),
        decision="rejected",
        codes=("QR_SIGNATURE_INVALID",),
    )
    add_punch(
        db_session,
        kiosk,
        None,
        age=timedelta(minutes=5),
        decision="needs_code",
        codes=("FACE_NOT_RECOGNISED",),
        method="face",
    )

    # Newest punch time first: the offline punch dates from when it was made (92 min ago).
    assert ids(boss.get(REVIEWS)) == [str(edge.id), str(travel.id)]
    assert boss.get(f"{REVIEWS}/summary").json()["to_check"] == 2
    day = to_ist(travel.punch_time).date().isoformat()
    log_summary = boss.get(f"{PUNCHES}/summary", params={"date": day}).json()
    assert log_summary["to_check"] == len(
        [p for p in (edge, travel) if to_ist(p.punch_time).date().isoformat() == day]
    )
    # Rejected: repeat scans hidden unless asked for; "asked for code" never.
    assert ids(boss.get(REVIEWS, params={"tab": "rejected"})) == [str(unknown.id), str(expired.id)]
    with_repeats = boss.get(REVIEWS, params={"tab": "rejected", "include_cooldown": True})
    assert ids(with_repeats) == [str(unknown.id), str(repeat.id), str(expired.id)]
    # Filters.
    assert ids(boss.get(REVIEWS, params={"code": "IMPOSSIBLE_TRAVEL"})) == [str(travel.id)]
    assert ids(boss.get(REVIEWS, params={"was_offline": True})) == [str(travel.id)]
    assert ids(boss.get(REVIEWS, params={"q": "asha"})) == [str(edge.id), str(travel.id)]
    named = boss.get(REVIEWS, params={"tab": "rejected", "names_employee": True})
    assert ids(named) == [str(expired.id)]
    nobody = boss.get(REVIEWS, params={"tab": "rejected", "names_employee": False})
    assert ids(nobody) == [str(unknown.id)]


def test_a_list_row_says_what_the_screen_needs(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    add_punch(
        db_session, kiosk, worker, age=timedelta(minutes=5), offline_delay=timedelta(minutes=52)
    )
    expired = add_punch(
        db_session,
        kiosk,
        worker,
        age=timedelta(minutes=1),
        decision="rejected",
        codes=("WRONG_PLANT",),
    )

    (row,) = boss.get(REVIEWS).json()["items"]
    assert row["punch"]["was_offline"] is True
    assert row["punch"]["offline_delay_seconds"] == 52 * 60
    assert row["punch"]["time_proof"] == "anchored"
    assert row["punch"]["review_state"] == "to_check"
    assert row["latest_review"] is None and row["approvable"] is False
    (rejected,) = boss.get(REVIEWS, params={"tab": "rejected"}).json()["items"]
    assert rejected["punch"]["id"] == str(expired.id) and rejected["approvable"] is True


def test_a_date_range_longer_than_62_days_is_refused(boss: TestClient) -> None:
    today = today_ist()
    response = boss.get(
        REVIEWS,
        params={
            "date_from": (today - timedelta(days=62)).isoformat(),
            "date_to": today.isoformat(),
        },
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "VALIDATION_FAILED"


# --- Roles (AC-29) -------------------------------------------------------------------------------


def test_review_endpoints_are_for_the_super_admin_only(
    corporate: TestClient,
    plant_head: TestClient,
    db_session: Session,
    kiosk: KioskDevice,
    worker: Worker,
) -> None:
    punch = add_punch(db_session, kiosk, worker, age=timedelta(minutes=5))

    for client in (corporate, plant_head):
        answers = [
            client.get(REVIEWS),
            client.get(f"{REVIEWS}/summary"),
            client.get(f"{REVIEWS}/digest"),
            client.get(f"{PUNCHES}/{punch.id}/review-context"),
            review(client, punch, "confirm"),
        ]
        for answer in answers:
            assert answer.status_code == 403, answer.text
            assert answer.json()["error"]["code"] == "FORBIDDEN"
    assert db_session.scalars(select(PunchReview)).all() == []


def test_other_roles_see_the_review_mark_as_text_without_the_reason(
    boss: TestClient,
    corporate: TestClient,
    plant_head: TestClient,
    db_session: Session,
    kiosk: KioskDevice,
    worker: Worker,
) -> None:
    punch = add_punch(db_session, kiosk, worker, age=timedelta(minutes=5))
    assert (
        review(boss, punch, "not_accepted", reason="Device was at the canteen").status_code == 201
    )
    day = to_ist(punch.punch_time).date().isoformat()

    for client in (corporate, plant_head):
        (item,) = client.get(PUNCHES, params={"date": day}).json()["items"]
        assert item["decision"] == "flagged" and item["review_state"] == "not_accepted"
        detail = client.get(f"{PUNCHES}/{punch.id}").json()
        assert detail["reviews"] == [
            {
                "id": None,
                "action": "not_accepted",
                "direction": None,
                "reason": None,
                "reviewed_by": None,
                "reviewed_at": None,
                "supersedes_review_id": None,
            }
        ]
        assert client.get(f"{PUNCHES}/summary", params={"date": day}).json()["to_check"] is None
    full = boss.get(f"{PUNCHES}/{punch.id}").json()["reviews"]
    assert full[0]["reason"] == "Device was at the canteen"
    assert full[0]["reviewed_by"]["name"]


def test_the_punch_log_filters_flagged_offline_and_review_state(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    flagged = add_punch(db_session, kiosk, worker, age=timedelta(minutes=3))
    offline = add_punch(
        db_session,
        kiosk,
        worker,
        age=timedelta(minutes=2),
        decision="accepted",
        offline_delay=timedelta(minutes=1),
    )
    confirmed = add_punch(db_session, kiosk, worker, age=timedelta(minutes=1))
    assert review(boss, confirmed, "confirm").status_code == 201
    day = to_ist(utc_now()).date().isoformat()

    def log_ids(**params: str | bool) -> set[str]:
        items = boss.get(PUNCHES, params={"date": day, **params}).json()["items"]
        return {item["id"] for item in items}

    if to_ist(flagged.punch_time).date().isoformat() != day:
        pytest.skip("the rows straddle IST midnight")
    assert log_ids(flagged=True) == {str(flagged.id), str(confirmed.id)}
    assert log_ids(decision="flagged") == {str(flagged.id), str(confirmed.id)}
    assert log_ids(was_offline=True) == {str(offline.id)}
    assert log_ids(review_state="to_check") == {str(flagged.id)}
    assert log_ids(review_state="confirmed") == {str(confirmed.id)}
    summary = boss.get(f"{PUNCHES}/summary", params={"date": day}).json()
    assert summary["flagged"] == 2 and summary["offline"] == 1 and summary["to_check"] == 1


def test_the_export_has_punch_time_offline_flags_and_review(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    punch = add_punch(
        db_session, kiosk, worker, age=timedelta(minutes=2), offline_delay=timedelta(minutes=1)
    )
    review(boss, punch, "confirm")
    day = to_ist(punch.punch_time).date().isoformat()

    response = boss.get(f"{PUNCHES}/export.xlsx", params={"date": day})

    book = openpyxl.load_workbook(io.BytesIO(response.content))
    header, row = [list(r) for r in book.active.iter_rows(values_only=True)]  # type: ignore[union-attr]
    values = dict(zip(header, row, strict=True))
    assert values["Result"] == "Flagged" and values["Reason code"] == ""
    assert values["Saved offline"] == "Yes"
    assert values["Flags"] == "At the edge of the fence (LOCATION_AT_EDGE)"
    assert values["Review"] == "Confirmed"
    assert values["Punch time"] == f"{to_ist(punch.punch_time):%d/%m/%Y %H:%M:%S}"


# --- Confirm and not accepted (AC-26) ------------------------------------------------------------


def test_not_accepted_stops_a_flagged_in_from_counting(
    boss: TestClient,
    db_client: TestClient,
    db_session: Session,
    device: Device,
    kiosk: KioskDevice,
    worker: Worker,
) -> None:
    flagged_in = add_punch(db_session, kiosk, worker, age=timedelta(minutes=60))

    response = review(boss, flagged_in, "not_accepted", reason="Device was taken outside")

    assert response.status_code == 201, response.text
    assert response.json()["review"]["action"] == "not_accepted"
    db_session.expire_all()
    stored = db_session.get(Punch, flagged_in.id)
    assert stored is not None and stored.decision == "flagged" and stored.direction == "in"
    (row,) = db_session.scalars(select(PunchReview)).all()
    assert row.employee_id == worker.employee_id and row.supersedes_review_id is None
    assert audit_actions(db_session) == ["punch_review.not_accepted"]
    # The next punch is decided as if that IN never counted: IN again, not OUT.
    assert send(db_client, device, worker.qr_text).json()["direction"] == "in"


def test_confirm_keeps_a_flagged_in_counting(
    boss: TestClient,
    db_client: TestClient,
    db_session: Session,
    device: Device,
    kiosk: KioskDevice,
    worker: Worker,
) -> None:
    flagged_in = add_punch(db_session, kiosk, worker, age=timedelta(minutes=60))

    assert review(boss, flagged_in, "confirm").status_code == 201

    assert audit_actions(db_session) == ["punch_review.confirm"]
    assert send(db_client, device, worker.qr_text).json()["direction"] == "out"


def test_a_reason_is_needed_and_checked_after_trimming(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    punch = add_punch(db_session, kiosk, worker, age=timedelta(minutes=5))

    for reason in (None, "  ab  ", "x" * 301):
        answer = review(boss, punch, "not_accepted", reason=reason)
        assert answer.status_code == 422 and answer.json()["error"]["code"] == "VALIDATION_FAILED"
    assert db_session.scalars(select(PunchReview)).all() == []


# --- Approve (AC-27) -----------------------------------------------------------------------------


def test_approving_a_rejected_code_punch_counts_it_with_the_chosen_direction(
    boss: TestClient,
    db_client: TestClient,
    db_session: Session,
    device: Device,
    kiosk: KioskDevice,
    worker: Worker,
) -> None:
    add_punch(db_session, kiosk, worker, age=timedelta(hours=3), decision="accepted")  # an IN
    spoof = add_punch(
        db_session,
        kiosk,
        None,
        age=timedelta(minutes=40),
        decision="rejected",
        codes=("SPOOF_SUSPECTED",),
        method="face",
        claimed=worker,
    )

    context = boss.get(f"{PUNCHES}/{spoof.id}/review-context").json()
    assert context["approvable"] is True
    assert context["suggestion"] == {"direction": "out"}
    assert context["claimed_employee"]["code"] == "T-1"
    response = review(boss, spoof, "approve", reason="Checked the photo", direction="out")

    assert response.status_code == 201, response.text
    (row,) = db_session.scalars(select(PunchReview)).all()
    assert row.direction == "out" and row.employee_id == worker.employee_id
    assert audit_actions(db_session) == ["punch_review.approve"]
    entry = db_session.scalars(
        select(AuditLog).where(AuditLog.action == "punch_review.approve")
    ).one()
    assert entry.after is not None and entry.after["direction"] == "out"
    # The approved OUT counts: the next punch is an IN.
    assert send(db_client, device, worker.qr_text).json()["direction"] == "in"
    day = to_ist(spoof.punch_time).date().isoformat()
    items = boss.get(PUNCHES, params={"date": day, "review_state": "approved"}).json()["items"]
    assert [(i["id"], i["review_direction"]) for i in items] == [(str(spoof.id), "out")]


def test_what_may_not_be_approved_or_confirmed(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    nobody = add_punch(
        db_session,
        kiosk,
        None,
        age=timedelta(minutes=9),
        decision="rejected",
        codes=("QR_SIGNATURE_INVALID",),
    )
    repeat = add_punch(
        db_session,
        kiosk,
        worker,
        age=timedelta(minutes=8),
        decision="rejected",
        codes=("COOLDOWN_ACTIVE",),
    )
    flagged = add_punch(db_session, kiosk, worker, age=timedelta(minutes=7))
    accepted = add_punch(db_session, kiosk, worker, age=timedelta(minutes=6), decision="accepted")
    asked = add_punch(
        db_session,
        kiosk,
        None,
        age=timedelta(minutes=5),
        decision="needs_code",
        codes=("FACE_NOT_RECOGNISED",),
        method="face",
    )
    expired = add_punch(
        db_session,
        kiosk,
        worker,
        age=timedelta(minutes=4),
        decision="rejected",
        codes=("WRONG_PLANT",),
    )

    def code(punch: Punch, action: str) -> str:
        answer = review(boss, punch, action, reason="Checked it", direction="in")
        assert answer.status_code == 409, answer.text
        return answer.json()["error"]["code"]

    assert code(nobody, "approve") == "REVIEW_NO_EMPLOYEE"
    assert code(repeat, "approve") == "REVIEW_NOT_ALLOWED"
    assert code(flagged, "approve") == "REVIEW_NOT_ALLOWED"
    assert code(accepted, "confirm") == "REVIEW_NOT_ALLOWED"
    assert code(asked, "confirm") == "REVIEW_NOT_ALLOWED"
    assert code(expired, "confirm") == "REVIEW_NOT_ALLOWED"
    assert code(expired, "not_accepted") == "REVIEW_NOT_ALLOWED"  # nothing to withdraw yet
    assert (
        review(boss, boss_punch := expired, "approve", "Sent to the other plant", "in").status_code
        == 201
    )
    assert boss_punch.id == expired.id
    assert db_session.scalars(select(PunchReview)).all() != []


def test_an_approval_can_be_withdrawn(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    expired = add_punch(
        db_session,
        kiosk,
        worker,
        age=timedelta(minutes=4),
        decision="rejected",
        codes=("WRONG_PLANT",),
    )
    first = review(boss, expired, "approve", "Sent to the other plant", "in").json()["review"]

    second = review(boss, expired, "not_accepted", "Wrong person", expected=first["id"])

    assert second.status_code == 201, second.text
    (item,) = boss.get(REVIEWS, params={"tab": "reviewed"}).json()["items"]
    assert item["punch"]["review_state"] == "not_accepted"


# --- Two Super Admins and the history (AC-28) ----------------------------------------------------


def test_a_second_review_of_what_was_shown_gets_review_changed(
    boss: TestClient,
    db_app: FastAPI,
    db_session: Session,
    kiosk: KioskDevice,
    worker: Worker,
) -> None:
    other_boss = client_for(db_app, db_session, Role.SUPER_ADMIN, "boss2")
    punch = add_punch(db_session, kiosk, worker, age=timedelta(minutes=5))

    first = review(boss, punch, "confirm")
    late = review(other_boss, punch, "not_accepted", reason="Device moved")

    assert first.status_code == 201
    assert late.status_code == 409 and late.json()["error"]["code"] == "REVIEW_CHANGED"
    assert late.json()["error"]["details"]["action"] == "confirmed"
    first_id = first.json()["review"]["id"]
    change = review(other_boss, punch, "not_accepted", reason="Device moved", expected=first_id)
    assert change.status_code == 201
    assert change.json()["review"]["supersedes_review_id"] == first_id
    history = boss.get(f"{PUNCHES}/{punch.id}/review-context").json()["reviews"]
    assert [h["action"] for h in history] == ["not_accepted", "confirmed"]
    assert audit_actions(db_session) == ["punch_review.confirm", "punch_review.not_accepted"]
    reviewed = boss.get(REVIEWS, params={"tab": "reviewed"}).json()["items"]
    assert [(r["punch"]["id"], r["latest_review"]["action"]) for r in reviewed] == [
        (str(punch.id), "not_accepted")
    ]
    assert boss.get(REVIEWS).json()["items"] == []  # nothing left to check


# --- The review panel (AC-25) --------------------------------------------------------------------


def test_the_panel_shows_the_fence_version_and_distance_of_an_edge_punch(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    fence = db_session.scalars(select(Geofence).where(Geofence.plant_id == kiosk.plant_id)).one()
    punch = add_punch(
        db_session,
        kiosk,
        worker,
        age=timedelta(minutes=5),
        distance_m=18.0,
        lat=GATE_LAT,
        lon=GATE_LON,
    )

    context = boss.get(f"{PUNCHES}/{punch.id}/review-context").json()

    assert context["geofence"]["version_id"] == str(fence.id)
    assert context["geofence"]["was_current_at_punch_time"] is True
    assert context["distance_m"] == 18.0 and context["margin_m"] == fence.buffer_m
    assert context["related"] == [] and context["suggestion"] is None
    assert context["punch"]["location"]["lat"] == pytest.approx(GATE_LAT)


def test_the_panel_shows_both_punches_of_impossible_travel(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker, k134: PlantIds
) -> None:
    far_gate = make_active_kiosk(db_session, k134, Device(), name="Far gate")
    earlier = add_punch(
        db_session,
        far_gate,
        worker,
        age=timedelta(minutes=12),
        decision="accepted",
        lat=GATE_LAT + 0.0216,
        lon=GATE_LON,  # about 2.4 km north
    )
    later = add_punch(
        db_session,
        kiosk,
        worker,
        age=timedelta(minutes=10),
        direction="out",
        codes=("IMPOSSIBLE_TRAVEL",),
        lat=GATE_LAT,
        lon=GATE_LON,
    )
    later_row = db_session.get(Punch, later.id)
    assert later_row is not None
    # The rule stored the link; rows are append-only, so the test builds a fresh linked row.
    linked = Punch(
        **{
            column.key: getattr(later_row, column.key)
            for column in Punch.__table__.columns
            if column.key not in ("id", "punch_time", "client_punch_id", "related_punch_id")
        },
        client_punch_id=uuid.uuid4(),
        related_punch_id=earlier.id,
    )
    db_session.add(linked)
    db_session.commit()

    context = boss.get(f"{PUNCHES}/{linked.id}/review-context").json()

    (related,) = context["related"]
    assert related["punch"]["id"] == str(earlier.id) and related["relation"] == "travel"
    assert related["distance_m"] == pytest.approx(2400, rel=0.02)
    assert related["minutes_apart"] == pytest.approx(2, abs=0.1)
    assert related["speed_kmh"] == pytest.approx(72, rel=0.03)
    assert context["limits"]["impossible_travel_kmh"] == 60


def test_the_approve_warning_lists_counted_punches_within_the_cooldown(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    counted = add_punch(db_session, kiosk, worker, age=timedelta(minutes=11), decision="accepted")
    expired = add_punch(
        db_session,
        kiosk,
        worker,
        age=timedelta(minutes=10),
        decision="rejected",
        codes=("WRONG_PLANT",),
    )

    context = boss.get(f"{PUNCHES}/{expired.id}/review-context").json()

    assert [n["punch_id"] for n in context["nearby_counted"]] == [str(counted.id)]


def test_an_unknown_punch_is_not_found(boss: TestClient) -> None:
    response = boss.get(f"{PUNCHES}/{uuid.uuid4()}/review-context")

    assert response.status_code == 404


# --- Hints (AC-43) -------------------------------------------------------------------------------


def test_hints_show_from_the_thresholds_on(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    for minutes in (50, 40):
        add_punch(
            db_session,
            kiosk,
            worker,
            age=timedelta(minutes=minutes),
            decision="rejected",
            codes=("WRONG_PLANT",),
        )
    third = add_punch(
        db_session,
        kiosk,
        worker,
        age=timedelta(minutes=30),
        decision="rejected",
        codes=("WRONG_PLANT",),
    )
    two = boss.get(f"{PUNCHES}/{third.id}/review-context").json()
    assert [h["kind"] for h in two["hints"]] == ["employee_rejections"]  # 3 = the threshold
    assert two["hints"][0]["count"] == 3
    assert two["hints"][0]["filter"]["employee_id"] == str(worker.employee_id)

    second = db_session.scalars(
        select(Punch).where(Punch.decision == "rejected").order_by(Punch.punch_time)
    ).all()[1]
    below = boss.get(f"{PUNCHES}/{second.id}/review-context").json()
    assert below["hints"] == []  # only 2 rejections up to that punch


def test_the_edge_flag_hint_needs_the_device_threshold(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    punches = [
        add_punch(db_session, kiosk, worker, age=timedelta(minutes=100 - 4 * n)) for n in range(20)
    ]

    hints = boss.get(f"{PUNCHES}/{punches[-1].id}/review-context").json()["hints"]
    kinds = {hint["kind"]: hint for hint in hints}
    assert kinds["device_edge_flags"]["count"] == 20
    assert kinds["device_edge_flags"]["filter"]["code"] == "LOCATION_AT_EDGE"
    nineteenth = boss.get(f"{PUNCHES}/{punches[-2].id}/review-context").json()["hints"]
    assert "device_edge_flags" not in {hint["kind"] for hint in nineteenth}


# --- Daily digest (AC-42) ------------------------------------------------------------------------


def test_the_digest_counts_yesterday_per_plant_with_offline_minutes_in_shifts(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    yesterday = today_ist() - timedelta(days=1)
    midnight = datetime.combine(yesterday, time.min, tzinfo=IST)
    hours_ago = (utc_now() - (midnight + timedelta(hours=9))).total_seconds() / 3600

    def at_nine(decision: str, codes: tuple[str, ...], who: Worker | None) -> Punch:
        return add_punch(
            db_session, kiosk, who, age=timedelta(hours=hours_ago), decision=decision, codes=codes
        )

    flagged = at_nine("flagged", ("LOCATION_AT_EDGE",), worker)
    at_nine("flagged", ("LOCATION_AT_EDGE",), worker)
    at_nine("rejected", ("WRONG_PLANT",), worker)
    at_nine("rejected", ("QR_SIGNATURE_INVALID",), None)  # names nobody
    at_nine("rejected", ("COOLDOWN_ACTIVE",), worker)  # nothing to approve
    review(boss, flagged, "confirm")
    # Health samples every 5 minutes from 06:00 to 20:00, with a silence from 10:00 to 12:08:
    # 128 minutes minus the 8 minutes that still count as "recently seen" = 2 h offline in shift.
    kiosk.registered_at = midnight - timedelta(days=1)
    samples = [midnight + timedelta(hours=6, minutes=5 * i) for i in range(49)]  # to 10:00
    samples += [midnight + timedelta(hours=12, minutes=8 + 5 * i) for i in range(95)]  # to 19:58
    for moment in samples:
        db_session.add(
            KioskHeartbeat(
                kiosk_id=kiosk.id, received_at=moment, queue_count=0, location_state="good"
            )
        )
    kiosk.last_seen_at = samples[-1] + timedelta(hours=1)
    db_session.commit()

    digest = boss.get(f"{REVIEWS}/digest").json()

    assert digest["date"] == yesterday.isoformat()
    (row,) = digest["plants"]
    assert row["plant"]["code"] == "K134"
    assert row["flagged"] == 2
    assert row["rejected_with_employee"] == 1
    assert row["to_check"] == 1
    assert row["device_offline_minutes_in_shifts"] == 120


def test_the_digest_goes_back_7_days_at_most(boss: TestClient) -> None:
    too_old = (today_ist() - timedelta(days=8)).isoformat()

    response = boss.get(f"{REVIEWS}/digest", params={"date": too_old})

    assert response.status_code == 422


# --- The audit log must not give back what the punch log hides (compliance review) ---


def test_the_audit_log_hides_review_reasons_from_the_other_roles(
    boss: TestClient,
    corporate: TestClient,
    plant_head: TestClient,
    db_session: Session,
    kiosk: KioskDevice,
    worker: Worker,
) -> None:
    punch = add_punch(db_session, kiosk, worker, age=timedelta(minutes=5))
    review(boss, punch, "not_accepted", reason="Used a friend's card, said the guard")

    def entry_of(client: TestClient) -> dict:
        items = client.get("/api/v1/audit", params={"action": "punch_review.not_accepted"}).json()
        (entry,) = items["items"]
        return entry

    full = entry_of(boss)
    assert full["reason"] == "Used a friend's card, said the guard"
    assert full["after"]["reason"] == "Used a friend's card, said the guard"
    for client in (corporate, plant_head):
        hidden = entry_of(client)
        assert hidden["reason"] is None and "reason" not in hidden["after"]
        assert "friend" not in str(hidden)
        assert hidden["action"] == "punch_review.not_accepted"  # that it happened stays visible


# --- An approved punch shows whom it counts for (compliance review, decision 0029) ---------------


def approved_code_punch(boss: TestClient, session: Session, kiosk: KioskDevice, worker: Worker):
    """A face punch refused by the face check, whose typed code names the worker, then approved."""
    punch = add_punch(
        session,
        kiosk,
        None,
        age=timedelta(minutes=5),
        decision="rejected",
        codes=("SPOOF_SUSPECTED",),
        method="face",
        claimed=worker,
    )
    assert review(boss, punch, "approve", "Checked the photo", "in").status_code == 201
    return punch


def test_an_approved_punch_without_an_employee_shows_whom_it_counts_for(
    boss: TestClient,
    plant_head: TestClient,
    corporate: TestClient,
    db_session: Session,
    kiosk: KioskDevice,
    worker: Worker,
) -> None:
    punch = approved_code_punch(boss, db_session, kiosk, worker)
    day = to_ist(punch.punch_time).date().isoformat()

    for client in (boss, corporate, plant_head):
        (item,) = client.get(PUNCHES, params={"date": day}).json()["items"]
        assert item["employee"] is None  # the punch row names nobody
        assert item["counted_for"]["code"] == "T-1" and item["review_state"] == "approved"
        detail = client.get(f"{PUNCHES}/{punch.id}").json()
        assert detail["counted_for"]["name"] == "Asha Kale"


def test_the_employee_filter_finds_the_punches_an_approval_counts_for_them(
    boss: TestClient,
    plant_head: TestClient,
    db_session: Session,
    kiosk: KioskDevice,
    worker: Worker,
    k134: PlantIds,
) -> None:
    punch = approved_code_punch(boss, db_session, kiosk, worker)
    day = to_ist(punch.punch_time).date().isoformat()
    other = Worker(boss, k134, "T-2")

    for client in (boss, plant_head):
        found = client.get(PUNCHES, params={"date": day, "employee_id": str(worker.employee_id)})
        assert [item["id"] for item in found.json()["items"]] == [str(punch.id)]
        nobody = client.get(PUNCHES, params={"date": day, "employee_id": str(other.employee_id)})
        assert nobody.json()["items"] == []


def test_a_plant_head_of_another_plant_never_learns_the_person(
    boss: TestClient,
    db_app: FastAPI,
    db_session: Session,
    kiosk: KioskDevice,
    worker: Worker,
) -> None:
    other_plant = create_plant(boss, db_session, "TP1")
    stranger = client_for(db_app, db_session, Role.PLANT_HEAD, "tp1head")
    user = db_session.scalars(select(User).where(User.username == "tp1head")).one()
    set_user_plant(db_session, user, uuid.UUID(other_plant.id))
    db_session.commit()
    punch = approved_code_punch(boss, db_session, kiosk, worker)

    # The punch was made at K134, so the TP1 head does not see it at all, and the person is hidden
    # from any filter by their id as well.
    day = to_ist(punch.punch_time).date().isoformat()
    assert stranger.get(PUNCHES, params={"date": day}).json()["items"] == []
    assert stranger.get(f"{PUNCHES}/{punch.id}").status_code == 404


def test_the_export_says_whom_an_approved_punch_counts_for(
    boss: TestClient, db_session: Session, kiosk: KioskDevice, worker: Worker
) -> None:
    punch = approved_code_punch(boss, db_session, kiosk, worker)
    day = to_ist(punch.punch_time).date().isoformat()

    response = boss.get(f"{PUNCHES}/export.xlsx", params={"date": day})

    book = openpyxl.load_workbook(io.BytesIO(response.content))
    header, row = [list(r) for r in book.active.iter_rows(values_only=True)]  # type: ignore[union-attr]
    values = dict(zip(header, row, strict=True))
    assert values["Counted for"] == "T-1 · Asha Kale"
    assert values["Review"] == "Approved as IN"

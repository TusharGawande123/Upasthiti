"""Excel import (all-or-nothing, row errors, injection, Corporate Head via approval) and the
contractors list (spec AC-22, AC-23, decision 0005)."""

import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.modules.audit.models import AuditLog
from app.modules.employees.import_excel import COLUMNS
from app.modules.employees.models import Employee
from app.modules.users.models import Role
from tests.people_helpers import PlantIds, client_for, create_plant

IMPORT = "/api/v1/employees/import"


def make_file(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(list(COLUMNS))
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def good_row(code: str, name: str, kind: str = "permanent", plant: str = "K134") -> list:
    rate = "85.00" if kind == "temporary" else ""
    return [code, name, kind, plant, "Staff" if kind == "permanent" else "Helper", "", "General",
            "", "", rate, "", "01/11/2026"]  # fmt: skip


def upload(client: TestClient, rows: list[list], dry_run: bool):
    files = {"file": ("people.xlsx", make_file(rows), "application/octet-stream")}
    return client.post(f"{IMPORT}?dry_run={str(dry_run).lower()}", files=files)


@pytest.fixture
def boss(db_app: FastAPI, db_session: Session) -> TestClient:
    return client_for(db_app, db_session, Role.SUPER_ADMIN, "boss")


@pytest.fixture
def k134(boss: TestClient, db_session: Session) -> PlantIds:
    return create_plant(boss, db_session, "K134")


def employee_count(session: Session) -> int:
    return len(session.scalars(select(Employee)).all())


def test_template_has_the_headers_and_two_example_rows(boss: TestClient, k134: PlantIds) -> None:
    response = boss.get("/api/v1/employees/import/template.xlsx")

    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content))["Employees"]
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == list(COLUMNS)
    assert len(rows) == 3  # header + two fake examples
    assert rows[1][3] == "K134"


def test_dry_run_lists_every_wrong_row_and_saves_nothing(
    boss: TestClient, k134: PlantIds, db_session: Session
) -> None:
    rows = [
        good_row("1001", "Ravi Patil"),
        good_row("1002", "Asha Rao", kind="temporary"),
        good_row("1001", "Duplicate code"),  # same code as row 2 of the data
        good_row("1003", "Wrong plant", plant="NOPE"),
        [
            "1004",
            "No rate",
            "temporary",
            "K134",
            "Helper",
            "",
            "General",
            "",
            "",
            "",
            "",
            "01/11/2026",
        ],
    ]

    response = upload(boss, rows, dry_run=True).json()

    assert response["ok_rows"] == 2
    assert response["error_rows"] == 3
    problems = {(error["row"], error["column"], error["code"]) for error in response["errors"]}
    assert (4, "employee_code", "EMPLOYEE_CODE_DUPLICATE_IN_FILE") in problems
    assert (5, "plant_code", "PLANT_NOT_FOUND") in problems
    assert (6, "hourly_rate", "HOURLY_RATE_REQUIRED") in problems
    assert employee_count(db_session) == 0


def test_saving_a_file_with_errors_saves_nothing_then_a_fixed_file_saves_all(
    boss: TestClient, k134: PlantIds, db_session: Session
) -> None:
    bad = [good_row("1001", "Ravi Patil"), good_row("1002", "Asha Rao", plant="NOPE")]

    refused = upload(boss, bad, dry_run=False)
    assert refused.status_code == 422
    assert refused.json()["error"]["code"] == "IMPORT_HAS_ERRORS"
    assert employee_count(db_session) == 0

    fixed = [good_row("1001", "Ravi Patil"), good_row("1002", "Asha Rao")]
    saved = upload(boss, fixed, dry_run=False).json()

    assert saved["created"] == 2
    assert employee_count(db_session) == 2
    entry = db_session.scalars(select(AuditLog).where(AuditLog.action == "employee.import")).one()
    assert entry.after == {"file_name": "people.xlsx", "rows": 2}


@pytest.mark.parametrize("cell", ["=1+1", "+cmd|' /C calc'!A0", "@SUM(A1)", "-2+3"])
def test_formula_like_text_is_refused(boss: TestClient, k134: PlantIds, cell: str) -> None:
    row = good_row("1001", cell)

    response = upload(boss, [row], dry_run=True).json()

    assert response["errors"][0]["code"] == "IMPORT_CELL_FORMULA"
    assert response["errors"][0]["column"] == "full_name"


def test_wrong_files_are_refused(boss: TestClient, k134: PlantIds) -> None:
    not_excel = boss.post(
        f"{IMPORT}?dry_run=true", files={"file": ("a.xlsx", b"not a workbook", "text/plain")}
    )
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["wrong", "headers"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    wrong_headers = boss.post(
        f"{IMPORT}?dry_run=true", files={"file": ("a.xlsx", buffer.getvalue(), "text/plain")}
    )

    assert not_excel.json()["error"]["code"] == "IMPORT_FILE_INVALID"
    assert wrong_headers.json()["error"]["code"] == "IMPORT_FILE_INVALID"


def test_corporate_head_import_waits_for_approval(
    boss: TestClient, k134: PlantIds, db_app: FastAPI, db_session: Session
) -> None:
    corporate = client_for(db_app, db_session, Role.CORPORATE_HEAD, "corp.head")
    rows = [good_row("1001", "Ravi Patil"), good_row("1002", "Asha Rao", kind="temporary")]

    sent = upload(corporate, rows, dry_run=False)

    assert sent.status_code == 202
    assert employee_count(db_session) == 0
    approved = boss.post(
        f"/api/v1/change-requests/{sent.json()['change_request_id']}/approve", json={}
    )
    assert approved.json()["status"] == "approved"
    assert employee_count(db_session) == 2


def test_plant_head_cannot_import(db_app: FastAPI, db_session: Session, k134: PlantIds) -> None:
    head = client_for(db_app, db_session, Role.PLANT_HEAD, "plant.head")

    assert upload(head, [good_row("1001", "Ravi Patil")], dry_run=True).json()["error"]["code"] == (
        "FORBIDDEN"
    )
    assert head.get("/api/v1/employees/import/template.xlsx").status_code == 403


def test_contractors_are_managed_by_the_super_admin_and_read_by_others(
    boss: TestClient, db_app: FastAPI, db_session: Session
) -> None:
    created = boss.post(
        "/api/v1/contractors",
        json={
            "name": "Shree Manpower",
            "contact_name": "Mr. Joshi",
            "contact_phone": "98 1234 5678",
        },
    )
    duplicate = boss.post("/api/v1/contractors", json={"name": "shree manpower"})
    head = client_for(db_app, db_session, Role.PLANT_HEAD, "plant.head")

    assert created.status_code == 201
    assert duplicate.json()["error"]["code"] == "CONTRACTOR_EXISTS"
    assert len(head.get("/api/v1/contractors").json()) == 1
    assert head.post("/api/v1/contractors", json={"name": "Other"}).json()["error"]["code"] == (
        "FORBIDDEN"
    )


def test_contractor_with_active_workers_cannot_be_switched_off(
    boss: TestClient, k134: PlantIds
) -> None:
    contractor = boss.post("/api/v1/contractors", json={"name": "Shree Manpower"}).json()
    body = {
        "employee_code": "T-1",
        "full_name": "Sunil More",
        "type": "temporary",
        "plant_id": k134.id,
        "category_id": k134.helper_id,
        "shift_id": k134.shift_id,
        "contractor_id": contractor["id"],
        "hourly_rate": "85.00",
        "joined_on": "2026-10-13",
    }
    boss.post("/api/v1/employees", json=body)

    response = boss.post(
        f"/api/v1/contractors/{contractor['id']}/deactivate", json={"reason": "No longer used"}
    )

    assert response.json()["error"]["code"] == "CONTRACTOR_IN_USE"
